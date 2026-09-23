from pathlib import Path
from tempfile import NamedTemporaryFile
import csv, re

GUIDE_WORDS=("guide","description","expected output","general description","specifications")

def _rows_openpyxl(path, sheet, n=20):
    import openpyxl
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws=wb[sheet]
    out=[]
    for row in ws.iter_rows(min_row=1,max_row=n,values_only=True): out.append(list(row))
    wb.close(); return out

def _score_header(row):
    vals=[str(v).strip() for v in row if v not in (None,"")]
    if not vals:return 0
    text=" ".join(vals).lower()
    hits=sum(k in text for k in ["country","pais","channel","category","categoria","brand","marca","sku","date","year","period","calendar","fact"])
    return hits*3+len(vals)

def workbook_snapshot(path,n=12):
    import openpyxl
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets=[]
    for ws in wb.worksheets:
        rows=[]
        for r in ws.iter_rows(min_row=1,max_row=n,values_only=True): rows.append(list(r))
        # Find likely header for a normal table; preserve first rows because AI needs mixed-header context.
        headers=[]
        if rows:
            idx=max(range(len(rows)), key=lambda i:_score_header(rows[i]))
            headers=[str(v).strip() if v not in (None,"") else f"Column_{j+1}" for j,v in enumerate(rows[idx])]
        sheets.append({"name":ws.title,"rows":rows,"columns":headers,"likely_header_row":(idx+1 if rows else 1)})
    wb.close(); return {"sheets":sheets}

def xlsx_to_csv(path, sheet, header_row=1):
    """Stream a workbook sheet into a temporary CSV so Polars remains the processing engine."""
    import openpyxl
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws=wb[sheet]
    tmp=NamedTemporaryFile(prefix="sdtp_",suffix=".csv",delete=False)
    with open(tmp.name,"w",newline="",encoding="utf-8-sig") as f:
        w=csv.writer(f)
        for i,row in enumerate(ws.iter_rows(values_only=True),start=1):
            if i < header_row: continue
            vals=list(row)
            if i==header_row:
                seen={}; clean=[]
                for j,v in enumerate(vals):
                    name=str(v).strip() if v not in (None,"") else f"Column_{j+1}"
                    seen[name]=seen.get(name,0)+1
                    clean.append(name if seen[name]==1 else f"{name}_{seen[name]}")
                w.writerow(clean)
            else:
                w.writerow(vals)
    wb.close(); return Path(tmp.name)

def read_xls_snapshot(path,n=12):
    import xlrd
    book=xlrd.open_workbook(path,on_demand=True)
    sheets=[]
    for sh in book.sheets():
        rows=[sh.row_values(i) for i in range(min(n,sh.nrows))]
        idx=max(range(len(rows)),key=lambda i:_score_header(rows[i])) if rows else 0
        headers=[str(v).strip() if v not in (None,"") else f"Column_{j+1}" for j,v in enumerate(rows[idx])] if rows else []
        sheets.append({"name":sh.name,"rows":rows,"columns":headers,"likely_header_row":idx+1})
    return {"sheets":sheets}
